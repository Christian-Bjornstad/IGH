from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from igh_merge.excel import HEADERS, ExcelReportWriter, local_relationship_comment
from igh_merge.models import MergedRow, SourceRow
from igh_merge.privacy import PrivacyError
from igh_merge.service import MergeService


def sample_row(percent: float = 2.5, mutation: float | None = 7.7) -> MergedRow:
    source = SourceRow(
        1,
        "ACGT" * 35,
        140,
        1000,
        "IGHV1-2_01",
        "IGHJ4_02",
        percent,
        percent,
        mutation,
        "Y",
        "Y",
        100.0,
        "CARDR",
    )
    return MergedRow(
        source, "SYN_A", 1, 123_456, "Leader", "gDNA", "2099_01_02",
        other_samples="SYN_B_1_FR1", comment=""
    )


@pytest.mark.parametrize(('source', 'expected'), [
    ('Supports Leader-1 | IMGT 3.8.3: productive; identity=98% | ARResT: unassigned', 'Støtter Leader-1'),
    ('(Supports Leader-3); Supports Leader-1 and Leader-2 | IMGT: not productive', '(Støtter Leader-3); Støtter Leader-1 og Leader-2'),
    ('Variant of Leader-2 | IMGT: productive', '(Variant av Leader-2)'),
    ('(Variant av leader 2); Støtter Leader-3', '(Variant av Leader-2); Støtter Leader-3'),
    ('IMGT: productive; J=IGHJ4*02 F | ARResT: note', ''),
])
def test_export_comments_keep_norwegian_relationships_only(source, expected):
    assert local_relationship_comment(source) == expected


def test_export_cleans_comments_without_mutating_analysis_or_other_values(tmp_path):
    row = sample_row()
    row.comment = 'Variant of Leader-1 | IMGT: productive | ARResT: unassigned'
    output = tmp_path / 'merged.xlsx'
    ExcelReportWriter().write([row], output)
    wb = load_workbook(output)
    try:
        assert wb.active['W2'].value == '(Variant av Leader-1)'
        assert wb.active['C2'].value == row.source.sequence
        assert wb.active['P2'].value == row.sample
        assert row.comment.endswith('ARResT: unassigned')
    finally:
        wb.close()


def test_excel_layout_formula_and_formats(tmp_path: Path):
    output = tmp_path / "merged.xlsx"
    ExcelReportWriter().write([sample_row()], output)

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Sheet1"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet["O2"].value == '=IF(J2="","",100-J2)'
    assert sheet["Q2"].value == 123_456
    assert sheet["X2"].value.startswith(">SYN_A-Leader-1\n")
    assert sheet["P2"].font.bold is True
    assert not sheet["X2"].alignment.wrap_text
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:X2"
    assert len(sheet.conditional_formatting) == 3
    assert workbook.calculation.fullCalcOnLoad is True
    workbook.close()


def test_blank_mutation_keeps_identity_formula(tmp_path: Path):
    output = tmp_path / "blank.xlsx"
    ExcelReportWriter().write([sample_row(mutation=None)], output)
    workbook = load_workbook(output, data_only=False)
    assert workbook["Sheet1"]["J2"].value is None
    assert workbook["Sheet1"]["O2"].value == '=IF(J2="","",100-J2)'
    workbook.close()


def test_selected_excel_rows_are_filled_yellow_across_all_columns(tmp_path: Path):
    output = tmp_path / "yellow.xlsx"
    rows = [sample_row(), sample_row(percent=1.0)]
    ExcelReportWriter().write(rows, output, highlight_rows={0})

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["Sheet1"]
    assert all(cell.fill.fill_type == "solid" for cell in sheet[2])
    assert all(cell.fill.fgColor.rgb.endswith("FFFF00") for cell in sheet[2])
    assert all(cell.fill.fill_type is None for cell in sheet[3])
    workbook.close()


def test_service_uses_same_functional_candidates_for_excel_highlight(tmp_path: Path):
    captured = {}

    class CapturingWriter:
        def write(self, rows, output, **options):
            captured.update(options)
            return output

    rows = (sample_row(percent=2.5), sample_row(percent=2.49))
    result = SimpleNamespace(rows=rows)
    output = tmp_path / "merged.xlsx"

    written = MergeService(writer=CapturingWriter()).export(
        result, output, highlight_functional_rows=True
    )

    assert written == output
    assert captured["highlight_rows"] == frozenset({0})


def test_existing_output_is_not_silently_overwritten(tmp_path: Path):
    output = tmp_path / "merged.xlsx"
    output.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        ExcelReportWriter().write([sample_row()], output)
    assert output.read_bytes() == b"existing"


def test_output_inside_git_tree_is_blocked(tmp_path: Path):
    (tmp_path / ".git").mkdir()
    with pytest.raises(PrivacyError):
        ExcelReportWriter().write([sample_row()], tmp_path / "clinical.xlsx")


@pytest.mark.parametrize("percent", [2.39, 2.40, 2.49, 2.50])
def test_percent_boundaries_remain_numeric(tmp_path: Path, percent: float):
    output = tmp_path / f"merged-{percent}.xlsx"
    ExcelReportWriter().write([sample_row(percent=percent)], output)
    workbook = load_workbook(output, data_only=False)
    assert workbook["Sheet1"]["H2"].value == percent
    workbook.close()


def test_excel_places_controls_first_in_required_order(tmp_path: Path):
    rows = [
        sample_row(),
        sample_row(),
        sample_row(),
        sample_row(),
        sample_row(),
    ]
    rows[0].sample = "SYN_PATIENT"
    rows[1].sample = "NK"
    rows[2].sample = "IGH_POS"
    rows[3].sample = "NGS_NEG"
    rows[4].sample = "IGH_SHM_POS"
    output = tmp_path / "ordered.xlsx"

    ExcelReportWriter().write(rows, output)

    workbook = load_workbook(output, data_only=False)
    assert [workbook["Sheet1"][f"P{row}"].value for row in range(2, 7)] == [
        "IGH_SHM_POS",
        "IGH_POS",
        "NGS_NEG",
        "NK",
        "SYN_PATIENT",
    ]
    workbook.close()
