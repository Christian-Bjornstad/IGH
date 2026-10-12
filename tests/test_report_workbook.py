from pathlib import Path

import pytest
from openpyxl import load_workbook

from igh_merge.report_workbook import write_sample_workbook


def record(label="Leader-1"):
    return {
        "rank_label": label, "external_id": "SEQ-ABC123", "number": "1",
        "fasta": ">SEQ-ABC123\nACGT", "sequence": "ACGT",
        "genes": "IGHV1-2*01 / IGHD1-1*01 / IGHJ4*02", "identity": "97.0 % (97/100 nt)",
        "functionality": "Non-functional (unproductive)", "cdr3": "CARDR",
        "cdr3_length": "5", "indel": "Not detected", "imgt_subset": "Not assigned",
        "arrest_subset": "ARResT analysis missing", "raw_identities": "97.0",
        "j_notes": "-", "observations": "Manual assessment", "analysis_status": "Available",
        "molecule_type": "gDNA", "comment": "", "depth_leader": "100000",
        "depth_fr1": "50000", "summary_text": "Sequence: SEQ-ABC123\nResult: unproductive",
        "missing_evidence": "IMGT evidence missing", "images": (), "arrest_notes": "",
    }


def test_sample_workbook_has_review_tabs_and_correct_sequence(tmp_path):
    path = tmp_path / "sample.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [record(), record("FR1-3")],
                          run_date="2099_01_02", imgt_version="3.8.3", imgt_release="209901-1")
    assert path.exists(), "The per-sample review workbook was not generated"
    with_workbook = load_workbook(path)
    try:
        assert with_workbook.sheetnames == ["Oversikt", "Leader-1", "FR1-3", "Mal for Sanger"]
        assert with_workbook["Oversikt"]["A1"].value == "SYN_SAMPLE"
        sheet = with_workbook["Leader-1"]
        assert sheet["A3"].value == ">SEQ-ABC123"
        assert sheet["A4"].value == "ACGT"
        values = [cell.value for row in sheet for cell in row]
        assert "Non-functional (unproductive)" in values
        assert "97.0 % (97/100 nt)" in values
        assert sheet["A7"].value == "Target fraction"
        assert sheet["A8"].value == "FR1 support"
        assert "Identity source" not in values
        assert not any(cell.fill.fgColor.rgb == "00FFFF00" for row in sheet for cell in row)
        assert any("not provided" in str(cell.value) for row in with_workbook["Oversikt"] for cell in row)
        assert with_workbook["Oversikt"].max_row == 8
        assert with_workbook["Oversikt"].freeze_panes is None
    finally:
        with_workbook.close()


def test_clinical_text_is_not_executed_as_excel_formula(tmp_path):
    rec = record()
    rec["comment"] = '=HYPERLINK("https://example.test", "open")'
    path = tmp_path / "sample.xlsx"
    write_sample_workbook(path, "=SYN_SAMPLE", [rec], run_date="2099_01_02",
                          imgt_version="", imgt_release="", lymphotrack_version="1.2.3")
    assert path.exists(), "The per-sample review workbook was not generated"
    wb = load_workbook(path, data_only=False)
    try:
        assert wb["Oversikt"]["A1"].data_type == "s"
        comments = [cell for row in wb["Leader-1"] for cell in row if cell.value == rec["comment"]]
        assert comments and all(cell.data_type == "s" for cell in comments)
        assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    finally:
        wb.close()


def test_sample_workbook_protects_existing_output_and_git_tree(tmp_path):
    path = tmp_path / "sample.xlsx"
    path.write_bytes(b"original")
    with pytest.raises(FileExistsError):
        write_sample_workbook(path, "SYN_SAMPLE", [], run_date="", imgt_version="", imgt_release="")
    assert path.read_bytes() == b"original"
    from igh_merge.privacy import PrivacyError
    root = tmp_path / "repo"
    root.mkdir()
    (root / ".git").mkdir()
    with pytest.raises(PrivacyError):
        write_sample_workbook(root / "sample.xlsx", "SYN_SAMPLE", [], run_date="", imgt_version="", imgt_release="")


def test_failed_publish_leaves_no_partial_workbook(tmp_path, monkeypatch):
    def fail(*args):
        raise OSError("disk unavailable")

    monkeypatch.setattr("igh_merge.report_workbook.os.link", fail)
    path = tmp_path / "sample.xlsx"
    with pytest.raises(OSError, match="disk unavailable"):
        write_sample_workbook(path, "SYN_SAMPLE", [record()], run_date="", imgt_version="", imgt_release="")
    assert not path.exists()
    assert not list(tmp_path.glob(".sample-*.xlsx"))


def test_summary_preserves_source_columns_beyond_four(tmp_path):
    rec = record()
    rec["summary_rows"] = (("CDR2-IMGT", "3", "no", "60", "98"),)
    path = tmp_path / "sample.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [rec], run_date="", imgt_version="", imgt_release="")
    wb = load_workbook(path)
    try:
        sheet = wb["Leader-1"]
        row = next(cell.row for row in sheet for cell in row if cell.value == "CDR2-IMGT")
        assert sheet.cell(row, 4).value == "60 | 98"
    finally:
        wb.close()


def test_clean_observations_bold_only_positive_results_and_keep_coordinates(tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    rec = record()
    rec["observations"] = (
        "J-TRP: påvist (AA position 118; manual assessment; IMGT 4_IMGT-gapped-AA-sequences.txt FR4-IMGT (AA 118–121)); "
        "J-PHE: ikke påvist (AA position 118; manual assessment; IMGT 4_IMGT-gapped-AA-sequences.txt FR4-IMGT (AA 118–121)); "
        "G-X-G: Påvist (AA positions 119–121)"
    )
    path = tmp_path / "observations.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [rec], run_date="", imgt_version="", imgt_release="")
    wb = load_workbook(path, rich_text=True)
    try:
        sheet = wb["Leader-1"]
        row = next(c.row for values in sheet for c in values if c.value == "Manual observations")
        value = sheet.cell(row, 2).value
        assert isinstance(value, CellRichText)
        assert "manual assessment" not in str(value)
        assert "4_IMGT-gapped" not in str(value)
        assert "FR4-IMGT (AA 118–121)" in str(value)
        assert ";\n" in str(value)
        bold = [part.text for part in value if isinstance(part, TextBlock) and part.font.b]
        assert bold == ["J-TRP: påvist", "G-X-G: Påvist"]
        assert sheet.row_dimensions[row].height >= 75
        assert sheet.cell(row, 2).alignment.vertical == "center"
    finally:
        wb.close()


def test_summary_matches_imgt_colors_without_duplicate_text_or_links(tmp_path):
    rec = record()
    rec["summary_rows"] = (
        ("Result summary: SEQ-ABC123", "Unproductive IGH rearranged sequence (stop codons)"),
        ("V-GENE and allele", "Homsap IGHV1-2*01 F", "score = 100", "identity = 97.0% (97/100 nt)"),
    )
    rec["leader_fraction"] = "12.34 %"
    rec["fr1_support"] = "10.23 %"
    rec["arrest_seqcure_details"] = "FR3-IMGT shows mutational load imbalance (60.00% ratio)"
    rec["arrest_seqcure_advice"] = "Analyze with insertion/deletion search."
    path = tmp_path / "styled.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [rec], run_date="", imgt_version="", imgt_release="")
    wb = load_workbook(path, rich_text=True)
    try:
        sheet = wb["Leader-1"]
        values = [c.value for row in sheet for c in row]
        assert sheet["B7"].value == "12.34 %"
        assert sheet["B8"].value == "10.23 %"
        assert "IMGT result fields" in values
        assert not any("copyable" in str(v) for v in values)
        assert "Specialist assessment (manual)" not in values
        assert "IMGT summary (copyable source text)" not in values
        assert "ARResT screenshot (add manually)" not in values
        assert rec["arrest_seqcure_details"] in values
        assert rec["arrest_seqcure_advice"] in values
        row = next(c.row for cells in sheet for c in cells if c.value == "V-GENE and allele")
        assert sheet.cell(row, 1).fill.fgColor.rgb[-6:] == "FEFFDF"
        assert sheet.cell(row, 2).font.color.rgb[-6:] == "170A9B"
        assert sheet.cell(row, 1).font.color.rgb[-6:] == "000000"
        from openpyxl.cell.rich_text import TextBlock
        identity = sheet.cell(row, 4).value
        orange = [part.text for part in identity if isinstance(part, TextBlock) and part.font.color.rgb[-6:] == "C75B00"]
        assert orange == ["97.0%"]
        assert all(c.hyperlink is None for cells in sheet for c in cells)
        assert sheet.row_dimensions[6].height == 15
    finally:
        wb.close()


def test_sanger_matches_the_supplied_blank_template(tmp_path):
    path = tmp_path / "sanger.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [], run_date="", imgt_version="", imgt_release="")
    wb = load_workbook(path)
    try:
        sheet = wb["Mal for Sanger"]
        assert sheet.column_dimensions["A"].width == 167.85546875
        assert not sheet.merged_cells.ranges
        assert sheet.freeze_panes is None
        assert sheet["A1"].value == "Resultater IGHV-SHM Sangersekvensering"
        assert sheet["A1"].font.sz == 22 and sheet["A1"].font.bold and sheet["A1"].font.u == "single"
        assert sheet["A37"].value == "Hvis DNA:"
        assert sheet["A39"].value == "Leader (M1): 500-570 bp (DNA) "
        assert sheet["A47"].value == "Leader (M1): 415-485 bp (RNA)"
        assert sheet["A59"].value == "Singelplex – Agilent"
        assert sheet["A88"].value == "IMGT Leader"
        assert sheet.row_dimensions[1].height == 28.5
        assert sheet.row_dimensions[37].height == 21
        assert sheet["A45"].alignment.vertical == "center"
        assert all(sheet.cell(row, 1).value is None for row in range(4, 37))
    finally:
        wb.close()


def test_images_use_native_crop_with_original_bytes_and_nonoverlapping_anchors(tmp_path):
    from PIL import Image as RasterImage, ImageDraw
    from zipfile import ZipFile
    from xml.etree import ElementTree
    source = tmp_path / "01_v_gene.png"
    image = RasterImage.new("RGB", (1600, 800), "white")
    ImageDraw.Draw(image).rectangle((30, 40, 700, 720), fill="black")
    image.save(source)
    next_source = tmp_path / "03_j_gene.png"
    image.save(next_source)
    rec = record()
    rec["images"] = (str(source), str(next_source))
    path = tmp_path / "images.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [rec], run_date="", imgt_version="", imgt_release="")
    with ZipFile(path) as package:
        assert package.read("xl/media/image1.png") == source.read_bytes()
        assert package.read("xl/media/image2.png") == next_source.read_bytes()
        drawing = ElementTree.fromstring(package.read("xl/drawings/drawing1.xml"))
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
              "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
        rectangles = drawing.findall(".//a:srcRect", ns)
        assert rectangles[0].attrib == {"l": "875", "t": "3000", "r": "55188", "b": "7875"}
        assert all(value.isdecimal() for rectangle in rectangles for value in rectangle.attrib.values())
        assert [props.attrib["id"] for props in drawing.findall(".//xdr:cNvPr", ns)] == ["1", "2"]
        assert [geometry.attrib["prst"] for geometry in drawing.findall(".//a:prstGeom", ns)] == ["rect", "rect"]
    wb = load_workbook(path)
    try:
        picture = wb["Leader-1"]._images[0]
        assert picture.anchor.pic.blipFill.srcRect.r == 55188
        assert picture.anchor.ext.cx > 740 * 9525
        next_picture = wb["Leader-1"]._images[1]
        height_in_rows = picture.anchor.ext.cy / 9525 / 20
        assert next_picture.anchor._from.row > picture.anchor._from.row + height_in_rows
    finally:
        wb.close()


def test_crop_ignores_only_decorative_edge_rules(tmp_path):
    from PIL import Image as RasterImage, ImageDraw
    from igh_merge.report_workbook import _image_box
    source = tmp_path / "01_v_gene.png"
    image = RasterImage.new("RGB", (1600, 800), "white")
    drawing = ImageDraw.Draw(image)
    drawing.line((0, 1, 1599, 1), fill=(204, 204, 204), width=2)
    drawing.rectangle((30, 40, 700, 720), fill="black")
    image.save(source)
    assert _image_box(source) == (14, 24, 717, 737)
    # A colored annotation on an edge row is content, even beside a gray rule.
    drawing.point((1300, 1), fill=(220, 100, 100))
    image.save(source)
    assert _image_box(source)[2] == 1600


def test_long_seqcure_text_is_complete_across_rows_below_excel_height_limit(tmp_path):
    rec = record()
    text = "SeqCure detailed advice: " + ("Keep every finding, including  double spaces and coordinates.\n" * 105)
    rec["arrest_seqcure_advice"] = text
    path = tmp_path / "long_advice.xlsx"
    write_sample_workbook(path, "SYN_SAMPLE", [rec], run_date="", imgt_version="", imgt_release="")
    wb = load_workbook(path)
    try:
        sheet = wb["Leader-1"]
        rows = [cell.row for values in sheet for cell in values if str(cell.value).startswith("SeqCure advice")]
        assert len(rows) > 1
        assert "".join(sheet.cell(row, 2).value for row in rows) == text
        assert all(sheet.row_dimensions[row].height <= 300 for row in rows)
        assert all(sheet.cell(row, 2).data_type == "s" for row in rows)
    finally:
        wb.close()
