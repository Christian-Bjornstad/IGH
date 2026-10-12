from dataclasses import replace
import json
import zipfile
import pytest
from openpyxl import load_workbook

from igh_merge.external import create_external_batch, ImgtBatchResult, parse_imgt_full_result
from igh_merge.reports import generate_clinical_report_package, _sample_report_records
from igh_merge.evidence import EvidenceManifest
from test_external import merged_row
from test_imgt_observations import full_files
from test_external import SEQUENCE
from igh_merge.evidence import save_evidence_manifest
from igh_merge.external import parse_arrest_result
from test_external import arrest_text, candidate_batch


def results(batch):
    records = []
    for c in batch.candidates:
        files = full_files()
        files['1_Summary.txt'] = files['1_Summary.txt'].replace('SEQ-ABC123', c.external_id).replace(SEQUENCE, c.sequence)
        files['5_AA-sequences.txt'] = files['5_AA-sequences.txt'].replace('SEQ-ABC123', c.external_id)
        records.extend(parse_imgt_full_result(files, replace(batch, candidates=(c,))).records)
    return ImgtBatchResult({'IMGT/V-QUEST program version': '3.8.3'}, tuple(records), '', '')


def test_report_target_and_exact_support():
    rows = [merged_row(), merged_row('TGCA' * 40), merged_row('TGCA' * 20)]
    rows[0].source = replace(rows[0].source, rank=1)
    rows[1].source = replace(rows[1].source, rank=10)
    rows[2].target = 'FR1'
    rows[2].comment = '(Supports Leader-10)'
    batch = create_external_batch('2099_01_02', list(enumerate(rows)))
    data = _sample_report_records(rows, batch, results(batch), None)['SYN_LOCAL_ONLY']
    assert data[0]['fr1_support'] == 'Not detected'
    assert [record['rank_label'] for record in data] == ['Leader-1', 'Leader-10']
    assert data[1]['fr1_support'] == '3.00 %'
    assert '99.2' in data[0]['identity']


def test_report_supports_unproductive_leader_and_preserves_selected_low_leader():
    leader = merged_row('TGCA' * 40)
    leader.source = replace(leader.source, rank=4, in_frame='N', no_stop_codon='N')
    fr1 = merged_row('TGCA' * 20)
    fr1.target = 'FR1'
    fr1.source = replace(fr1.source, rank=2, in_frame='N', no_stop_codon='N')
    low_leader = merged_row('ACGT' * 40)
    low_leader.source = replace(low_leader.source, rank=7, percent_total_reads=1.0)
    orphan = merged_row('GATTACA' * 12)
    orphan.target = 'FR1'
    orphan.source = replace(orphan.source, rank=3)
    rows = [orphan, low_leader, fr1, leader]
    batch = create_external_batch('2099_01_02', list(enumerate(rows)))
    data = _sample_report_records(rows, batch, results(batch), None)['SYN_LOCAL_ONLY']
    assert [record['rank_label'] for record in data] == ['Leader-4', 'Leader-7', 'FR1-3']
    assert data[0]['fr1_support'] == '3.00 %'
    assert data[0]['supporting_external_ids'] == [fr1.external_id]
    assert 'manually selected' in data[1]['comment'].lower()
    assert data[0]['sequence'] == leader.source.sequence
    assert data[0]['fasta'] == '>SYN_LOCAL_ONLY-Leader-4\n' + leader.source.sequence


def test_report_does_not_hide_fr1_supporting_only_unselected_leader():
    rows = [merged_row('TGCA' * 40), merged_row('TGCA' * 20)]
    rows[1].target = 'FR1'
    rows[1].comment = 'Supports Leader-1'
    batch = create_external_batch('2099_01_02', [(1, rows[1])])
    data = _sample_report_records(rows, batch, results(batch), None)['SYN_LOCAL_ONLY']
    assert [record['rank_label'] for record in data] == ['FR1-1']


def test_report_keeps_fr1_when_selected_leader_is_below_threshold():
    rows = [merged_row('TGCA' * 40), merged_row('TGCA' * 20)]
    rows[0].source = replace(rows[0].source, percent_total_reads=2.49)
    rows[1].target = 'FR1'
    batch = create_external_batch('2099_01_02', list(enumerate(rows)))
    data = _sample_report_records(rows, batch, results(batch), None)['SYN_LOCAL_ONLY']
    assert [record['rank_label'] for record in data] == ['Leader-1', 'FR1-1']


@pytest.mark.parametrize('changes', [
    {'sequence': 'TGCA' * 40},
    {'molecule_type': 'cDNA'},
])
def test_report_rejects_candidate_payload_disagreeing_with_original_row(changes):
    batch = candidate_batch()
    imgt = results(batch)
    changed = replace(batch, candidates=(replace(batch.candidates[0], **changes),))
    with pytest.raises(ValueError, match='Report row does not match'):
        _sample_report_records([merged_row()], changed, imgt, None)


def test_package_audits_reported_and_supporting_sequences(tmp_path):
    rows = [merged_row('TGCA' * 40), merged_row('TGCA' * 20)]
    rows[1].target = 'FR1'
    batch = create_external_batch('2099_01_02', list(enumerate(rows)))
    directory = generate_clinical_report_package(tmp_path, rows, batch, results(batch))
    audit = json.loads((directory / 'report_data.audit.json').read_text(encoding='utf-8'))
    assert audit['reported_ids'] == [batch.candidate_ids[0]]
    assert audit['supporting_ids'] == [batch.candidate_ids[1]]
    assert audit['supporting_by_reported_id'] == {batch.candidate_ids[0]: [batch.candidate_ids[1]]}
    assert len(audit['mapping']) == len(audit['imgt_records']) == 2
    workbook = load_workbook(directory / 'analysis.xlsx')
    assert {workbook.active.cell(i, 1).value for i in (2, 3)} == set(batch.candidate_ids)
    assert (directory / 'SYN_LOCAL_ONLY_IGHV_analysis.xlsx').exists()


def test_report_preserves_arrest_source_details():
    batch = candidate_batch()
    imgt = results(batch)
    arrest = parse_arrest_result(arrest_text(), batch)
    record = _sample_report_records([merged_row()], batch, imgt, arrest)['SYN_LOCAL_ONLY'][0]
    assert record['arrest_metadata']['external_id'] == 'SEQ-ABC123'
    assert record['arrest_metadata']['subset'] == arrest.records[0].subset
    assert arrest.records[0].seqcure in record['arrest_notes']
    missing = _sample_report_records([merged_row()], batch, imgt, None)['SYN_LOCAL_ONLY'][0]
    assert missing['arrest_subset'] == 'ARResT analysis missing'
    assert missing['arrest_metadata'] == {}


def test_report_distinguishes_raw_cdr3_length_from_returned_aa_length():
    batch = candidate_batch()
    files = full_files(productive='unproductive')
    header, values = files['1_Summary.txt'].splitlines()
    files['1_Summary.txt'] = header + '\tCDR3-IMGT length\n' + values + '\tX'
    imgt = parse_imgt_full_result(files, batch)
    record = _sample_report_records([merged_row()], batch, imgt, None)['SYN_LOCAL_ONLY'][0]
    assert record['cdr3_length'] == 'IMGT: X; returned AA: 5'
    assert record['raw_cdr3_imgt_length'] == 'X'
    assert record['cdr3_aa_length'] == 5


def test_report_audit_exposes_separate_analysis_completeness(tmp_path):
    batch = candidate_batch()
    directory = generate_clinical_report_package(tmp_path, [merged_row()], batch, results(batch))
    audit = json.loads((directory / 'report_data.audit.json').read_text(encoding='utf-8'))
    assert audit['complete'] is audit['imgt_complete'] is True
    assert audit['arrest_complete'] is False
    assert audit['missing_arrest_ids'] == ['SEQ-ABC123']
    assert audit['evidence_complete'] is False
    assert audit['missing_evidence_ids'] == ['SEQ-ABC123']


def test_packaged_evidence_paths_are_relative_without_mutating_source(tmp_path):
    import base64
    import hashlib

    batch = candidate_batch()
    source = tmp_path / 'source'
    source.mkdir()
    image = source / '01_v_gene.png'
    image.write_bytes(base64.b64decode(
        'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jXioAAAAASUVORK5CYII='
    ))
    evidence = save_evidence_manifest(source, batch, batch.candidates[0],
                                      {'00_summary': 'Result summary: SEQ-ABC123'}, [image], parameters={})
    original = evidence.entries[0]
    directory = generate_clinical_report_package(tmp_path, [merged_row()], batch, results(batch), evidence=evidence)
    packaged = directory / 'evidence' / 'SEQ-ABC123'
    metadata = json.loads((packaged / 'evidence.audit.json').read_text(encoding='utf-8'))
    assert metadata['images'] == ['01_v_gene.png']
    assert metadata['image_hashes'] == {'01_v_gene.png': hashlib.sha256(image.read_bytes()).hexdigest()}
    assert (packaged / metadata['images'][0]).read_bytes() == image.read_bytes()
    assert original.images == (str(image.resolve()),)
    assert str(image.resolve()) in original.image_hashes


def test_package_excel_and_unique_sample_names(tmp_path):
    rows = [merged_row(), merged_row()]
    rows[0].sample, rows[1].sample = 'SYN/A', 'SYN?A'
    batch = create_external_batch('2099_01_02', list(enumerate(rows)))
    directory = generate_clinical_report_package(tmp_path, rows, batch, results(batch), evidence=EvidenceManifest(()))
    assert len(list(directory.glob('*.docx'))) == 2
    assert not list(directory.glob('*.pdf'))
    assert len(list(directory.glob('*_IGHV_analysis.xlsx'))) == 2
    workbook = load_workbook(directory / 'analysis.xlsx')
    assert workbook.active.max_column == 24
    assert {workbook.active.cell(i, 1).value for i in (2, 3)} == set(batch.candidate_ids)
    with zipfile.ZipFile(next(directory.glob('*.docx'))) as z:
        text = z.read('word/document.xml').decode()
        assert 'manual assessment' in text
        assert 'IMGT evidence missing' in text
    again = generate_clinical_report_package(tmp_path, rows, batch, results(batch))
    assert again != directory


def test_report_partial_results_mark_missing(tmp_path):
    rows = [merged_row(), merged_row()]
    batch = create_external_batch('2099_01_02', list(enumerate(rows)))
    imgt = replace(results(batch), records=results(batch).records[:1])
    directory = generate_clinical_report_package(tmp_path, rows, batch, imgt)
    audit = json.loads((directory / 'report_data.audit.json').read_text())
    assert audit['complete'] is False
    assert audit['missing_imgt_ids'] == [batch.candidate_ids[1]]


def test_missing_imgt_result_does_not_report_a_negative_indel_result():
    batch = candidate_batch()
    imgt = replace(results(batch), records=())
    record = _sample_report_records([merged_row()], batch, imgt, None)['SYN_LOCAL_ONLY'][0]
    assert record['indel'] == 'Not analyzed / IMGT result missing'


def test_evidence_coverage_does_not_imply_complete_required_sections(tmp_path):
    from igh_merge.evidence import EvidenceEntry

    batch = candidate_batch()
    candidate = batch.candidates[0]
    entry = EvidenceEntry(candidate.external_id, candidate.sequence_sha256, batch.session_id,
                          '2099-01-02T00:00:00Z', {}, {'00_summary': 'Captured source text'}, (), {}, {})
    directory = generate_clinical_report_package(tmp_path, [merged_row()], batch, results(batch),
                                                 evidence=EvidenceManifest((entry,)))
    audit = json.loads((directory / 'report_data.audit.json').read_text(encoding='utf-8'))
    assert audit['evidence_coverage_complete'] is True
    assert audit['missing_evidence_ids'] == []
    assert audit['evidence_complete'] is False
    assert set(audit['missing_evidence_sections_by_id']['SEQ-ABC123']) == {
        '00_summary', '01_v_gene', '03_j_gene', '06_junction', '09_v_region_translation',
    }
    assert not {'02_d_gene', '04_leader', '05_constant'} & set(audit['missing_evidence_sections_by_id']['SEQ-ABC123'])
