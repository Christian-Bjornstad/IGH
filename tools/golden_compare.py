#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


def _normalized(value):
    if isinstance(value, str):
        compact = value.replace("\u00a0", " ").strip()
        if compact.replace(" ", "").isdigit():
            return int(compact.replace(" ", ""))
        numeric = compact.replace(" ", "")
        if re.fullmatch(r"-?\d+(?:[,.]\d+)", numeric):
            return round(float(numeric.replace(",", ".")), 8)
        return compact.replace("\r\n", "\n")
    if isinstance(value, float):
        return round(value, 8)
    return value


def compare(expected: Path, actual: Path) -> list[str]:
    old = load_workbook(expected, data_only=True, read_only=True)
    new = load_workbook(actual, data_only=False, read_only=True)
    try:
        old_ws = old[old.sheetnames[0]]
        new_ws = new[new.sheetnames[0]]
        errors: list[str] = []
        if old_ws.max_row != new_ws.max_row or old_ws.max_column != new_ws.max_column:
            errors.append(
                f"Størrelse avviker: {old_ws.max_row}x{old_ws.max_column} mot "
                f"{new_ws.max_row}x{new_ws.max_column}"
            )
            return errors
        # N (identitet) og U (ny FR1-støtte) er dokumenterte forbedringer.
        compared_columns = [*range(1, 14), 15, 16, 17, 18, 19, 20, 22]
        old_rows = old_ws.iter_rows(values_only=True)
        new_rows = new_ws.iter_rows(values_only=True)
        for row, (old_values, new_values) in enumerate(zip(old_rows, new_rows), start=1):
            for column in compared_columns:
                left = _normalized(old_values[column - 1])
                right = _normalized(new_values[column - 1])
                if left != right:
                    errors.append(f"Avvik rad {row}, kolonne {column}: {left!r} != {right!r}")
                    if len(errors) >= 25:
                        return errors
        return errors
    finally:
        old.close()
        new.close()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("expected", type=Path)
    parser.add_argument("actual", type=Path)
    args = parser.parse_args()
    errors = compare(args.expected, args.actual)
    if errors:
        print("\n".join(errors))
        return 1
    print("Golden-sammenligning OK.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
