"""Per-sample review workbooks, separate from the compatible merged sheet."""

import os
import math
from pathlib import Path
import re
import tempfile
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.fill import Blip, RelativeRect
from openpyxl.drawing.geometry import PresetGeometry2D
from openpyxl.drawing.image import Image
from openpyxl.drawing.picture import PictureFrame
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as RasterImage, ImageChops

from .privacy import ensure_outside_git


INK = "172B3A"
MUTED = "526371"
SECTION = "E8EFF3"
BORDER = "CFD9DF"
IMGT_YELLOW = "FEFFDF"
IMGT_BLUE = "170A9B"
IMGT_ORANGE = "C75B00"
IMAGE_SECTIONS = {"01_v_gene", "02_d_gene", "03_j_gene", "06_junction", "09_v_region_translation"}
RESULT_FIELDS = (
    ("External ID", "external_id"), ("Target fraction", "leader_fraction"),
    ("FR1 support", "fr1_support"), ("IGHV / IGHD / IGHJ", "genes"),
    ("VH identity", "identity"), ("Functionality", "functionality"),
    ("CDR3 (AA)", "cdr3"), ("CDR3 length", "cdr3_length"),
    ("Insertion / deletion", "indel"), ("IMGT subset", "imgt_subset"),
    ("ARResT subset", "arrest_subset"), ("Raw identities", "raw_identities"),
    ("J call notes", "j_notes"), ("Manual observations", "observations"),
    ("Analysis status", "analysis_status"), ("Molecule type", "molecule_type"),
    ("Comment", "comment"),
)


def _text(sheet, row, column, value, *, bold=False):
    cell = sheet.cell(row, column)
    cell.value = value if isinstance(value, CellRichText) else str(value) if value is not None else ""
    # Clinical source text must remain literal, even if it starts with '='.
    cell.data_type = "s"
    cell.font = Font(name="Calibri", size=11, color=INK, bold=bold)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    return cell


def _wide_text(sheet, row, value, *, bold=False):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    _text(sheet, row, 1, value, bold=bold)


def _heading(sheet, row, title):
    _wide_text(sheet, row, title, bold=True)
    sheet.cell(row, 1).fill = PatternFill("solid", fgColor=SECTION)
    sheet.row_dimensions[row].height = 15


def _field(sheet, row, label, value):
    _text(sheet, row, 1, label, bold=True)
    sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    _text(sheet, row, 2, value or "-")
    sheet.cell(row, 1).fill = PatternFill("solid", fgColor=SECTION)
    for col in range(1, 5):
        sheet.cell(row, col).border = Border(bottom=Side(style="hair", color=BORDER))
    sheet.row_dimensions[row].height = 15 * max(_lines(label, 27), _lines(value or "-", 80))


def _lines(value, width):
    return sum(max(1, math.ceil(len(line) / width)) for line in str(value).split("\n"))


def _long_field(sheet, row, label, value):
    """Keep long source reports readable below Excel's 409-point row limit."""
    text = str(value)
    cursor = 0
    while cursor < len(text):
        limit = min(len(text), cursor + 1600)
        low, high = cursor + 1, limit
        while low < high:
            middle = (low + high + 1) // 2
            if _lines(text[cursor:middle], 80) <= 20:
                low = middle
            else:
                high = middle - 1
        end = low
        if end < len(text):
            boundary = max(text.rfind("\n", cursor, end), text.rfind(" ", cursor, end))
            if boundary >= max(cursor + 1, end - 80):
                end = boundary + 1
        _field(sheet, row, label if cursor == 0 else f"{label} (continued)", text[cursor:end])
        cursor = end
        if cursor < len(text):
            row += 1
    return row


def _configure(sheet):
    for col, width in {"A": 28, "B": 42, "C": 20, "D": 20}.items():
        sheet.column_dimensions[col].width = width
    sheet.sheet_format.defaultRowHeight = 15
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.3
    sheet.page_margins.top = sheet.page_margins.bottom = 0.4


def _overview(sheet, sample, records, run_date, imgt_version, imgt_release, lymphotrack_version):
    _wide_text(sheet, 1, sample, bold=True)
    sheet.cell(1, 1).font = Font(name="Calibri", size=18, color=INK, bold=True)
    sheet.row_dimensions[1].height = 30
    first = records[0] if records else {}
    fields = (
        ("Run date", run_date), ("Analysis", "LymphoTrack IGH SHM (HTS)"),
        ("LymphoTrack version", lymphotrack_version or "not provided"),
        ("IMGT/V-QUEST version", imgt_version or "not provided"),
        ("IMGT reference release", imgt_release or "not provided"),
        ("Depth", f"{first.get('depth_leader', '-')} reads (Leader); {first.get('depth_fr1', '-')} reads (FR1)"),
    )
    for row, (label, value) in enumerate(fields, start=3):
        _field(sheet, row, label, value)


def _observations(value):
    cleaned = re.sub(r";\s*manual assessment\s*;\s*IMGT\s+4_IMGT-gapped-AA-sequences\.txt\s*", "; ", str(value), flags=re.I)
    cleaned = re.sub(r";\s*", ";\n", cleaned)
    return _highlight(cleaned, r"(?m)(?:^|(?<=\n))[^;\n]*?:\s*påvist\b", INK)


def _highlight(value, pattern, color):
    pieces = []
    previous = 0
    for match in re.finditer(pattern, value, flags=re.I):
        if match.start() > previous:
            pieces.append(value[previous:match.start()])
        pieces.append(TextBlock(InlineFont(rFont="Calibri", sz=11, color=color, b=True), match.group()))
        previous = match.end()
    if not pieces:
        return value
    if previous < len(value):
        pieces.append(value[previous:])
    return CellRichText(pieces)


def _summary(sheet, row, values):
    label = str(values[0])
    visible = list(values[:3]) + ([" | ".join(str(value) for value in values[3:])] if len(values) > 3 else [])
    if len(values) == 2:
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    for column, value in enumerate(visible, start=1):
        value = str(value)
        color = "000000"
        if (column == 2 and "GENE and allele" in label) or (column == 3 and "decryption" in label):
            color = IMGT_BLUE
        elif label.startswith("Result summary"):
            pattern = r"(?<=Result summary: ).+" if column == 1 else r"(?:Unproductive|Productive) IGH rearranged sequence"
            value = _highlight(value, pattern, IMGT_ORANGE)
        elif column >= 2 and "identity" in value.lower():
            value = _highlight(value, r"\d+(?:\.\d+)?%", IMGT_ORANGE)
        cell = _text(sheet, row, column, value)
        cell.font = Font(name="Calibri", size=11, color=color)
    for col in range(1, 5):
        cell = sheet.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=IMGT_YELLOW)
        cell.border = Border(left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
                             top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"))
    widths = (27, 80) if len(values) == 2 else (27, 41, 19, 19)
    sheet.row_dimensions[row].height = 15 * max(_lines(v, widths[col]) for col, v in enumerate(visible))


def _image_box(filename):
    """Find the display rectangle, preserving every source byte in the workbook.

    IMGT captures contain decorative light-gray rules at their top/bottom edges.
    Ignoring only those edge pixels allows the large white right margin to be
    hidden by native Excel cropping; analysis text and colored pixels remain.
    """
    with RasterImage.open(filename) as original:
        rgb = original.convert("RGB")
        difference = ImageChops.difference(rgb, RasterImage.new("RGB", rgb.size, "white"))
        channels = difference.split()
        mask = ImageChops.lighter(ImageChops.lighter(channels[0], channels[1]), channels[2])
        pixels = rgb.load()
        selection = mask.load()
        for y in list(range(min(32, rgb.height))) + list(range(max(32, rgb.height - 32), rgb.height)):
            # A long, neutral rule is decorative; shorter marks or a row with
            # any dark/colored content may contain text and must stay intact.
            neutral = {x for x in range(rgb.width) if 180 <= pixels[x, y][0] < 255 and pixels[x, y][0] == pixels[x, y][1] == pixels[x, y][2]}
            if len(neutral) < rgb.width / 4 or any(selection[x, y] and x not in neutral for x in range(rgb.width)):
                continue
            for x in neutral:
                selection[x, y] = 0
        box = mask.getbbox()
        if not box:
            return (0, 0, rgb.width, rgb.height)
        left, top, right, bottom = box
        return (max(0, left - 16), max(0, top - 16), min(rgb.width, right + 16), min(rgb.height, bottom + 16))


def _picture(sheet, row, filename, image_id):
    picture = Image(filename)
    source_width, source_height = picture.width, picture.height
    left, top, right, bottom = _image_box(filename)
    visible_width, visible_height = right - left, bottom - top
    scale = 790 / visible_width
    picture.width, picture.height = visible_width * scale, visible_height * scale
    frame = PictureFrame()
    frame.nvPicPr.cNvPr.id = image_id
    frame.nvPicPr.cNvPr.name = Path(filename).stem
    frame.nvPicPr.cNvPr.descr = "Picture"
    frame.blipFill.blip = Blip(cstate="print")
    # DrawingML percentages are integer thousandths of a percent: 100000=100%.
    frame.blipFill.srcRect = RelativeRect(l=round(100000 * left / source_width), t=round(100000 * top / source_height),
                                         r=round(100000 * (source_width - right) / source_width),
                                         b=round(100000 * (source_height - bottom) / source_height))
    frame.spPr.prstGeom = PresetGeometry2D(prst="rect")
    frame.spPr.ln = None
    picture.anchor = OneCellAnchor(_from=AnchorMarker(col=0, row=row - 1),
                                  ext=XDRPositiveSize2D(pixels_to_EMU(picture.width), pixels_to_EMU(picture.height)), pic=frame)
    sheet.add_image(picture)
    occupied = math.ceil(picture.height / 20) + 1
    for picture_row in range(row, row + occupied):
        sheet.row_dimensions[picture_row].height = 15
    return row + occupied + 1


def _rearrangement(sheet, sample, rec):
    _heading(sheet, 1, f"Rearrangement {rec['number']} - {rec['rank_label']}")
    _wide_text(sheet, 2, f"Sample: {sample}")
    fasta = rec.get("fasta", f">{rec['external_id']}\n{rec.get('sequence', '')}")
    header, _, sequence = fasta.partition("\n")
    _wide_text(sheet, 3, header)
    _wide_text(sheet, 4, sequence or rec.get("sequence", ""))
    sheet.row_dimensions[4].height = max(30, 15 * (len(sequence) // 100 + 1))
    for row, (label, key) in enumerate(RESULT_FIELDS, start=6):
        value = rec.get(key, "-")
        _field(sheet, row, label, _observations(value) if key == "observations" else value)
    row = 6 + len(RESULT_FIELDS) + 1
    if rec.get("summary_rows"):
        _heading(sheet, row, "IMGT result fields")
        for values in rec["summary_rows"]:
            row += 1
            _summary(sheet, row, values)
        row += 2
    else:
        _heading(sheet, row, "IMGT result fields")
        for line in str(rec.get("summary_text", "IMGT evidence missing")).splitlines():
            row += 1
            _wide_text(sheet, row, line)
            sheet.row_dimensions[row].height = 15 * _lines(line, 100)
        row += 2
    _field(sheet, row, "ARResT notes", rec.get("arrest_notes") or "No notes provided")
    for label, key in (("SeqCure report", "arrest_seqcure_details"), ("SeqCure advice", "arrest_seqcure_advice"),
                       ("SeqCure details status", "arrest_seqcure_status")):
        if not rec.get(key):
            continue
        row += 1
        row = _long_field(sheet, row, label, rec[key])
    row += 3
    for image_id, filename in enumerate(rec.get("images", ()), start=1):
        if Path(filename).stem not in IMAGE_SECTIONS:
            continue
        _heading(sheet, row, Path(filename).stem.replace("_", " "))
        row = _picture(sheet, row + 1, filename, image_id)
    sheet.print_area = f"A1:D{row}"


def _sanger(sheet, sample):
    # Only the reusable labels/layout from the supplied blank template belong
    # here. Patient values and yellow reminder notes are deliberately excluded.
    labels = {
        1: "Resultater IGHV-SHM Sangersekvensering", 3: "Fragmentkjøring Identiclone",
        37: "Hvis DNA:", 38: "Referanseområder IGHV-hypermutasjon:", 39: "Leader (M1): 500-570 bp (DNA) ",
        40: "FR1 (M2): 310-380 bp", 41: "Resultat:", 42: "Miks 1: ", 43: "Miks 2: ",
        45: "Hvis RNA:", 46: "Referanseområder IGHV-hypermutasjon:", 47: "Leader (M1): 415-485 bp (RNA)",
        48: "FR1 (M2): 310-380 bp", 49: "Resultat:", 50: "Miks 1: ", 51: "Miks 2: ",
        53: "Nukleotidsekvenser Identiclone (FASTA)", 59: "Singelplex – Agilent",
        83: "Nukleotidsekvenser singleplex (FASTA)", 88: "IMGT Leader",
    }
    sheet.column_dimensions["A"].width = 167.85546875
    sheet.sheet_format.defaultRowHeight = 15
    for row, value in labels.items():
        size = 22 if row == 1 else 16 if row in (37, 45) else 15 if row in (3, 53, 59, 83, 88) else 11
        cell = _text(sheet, row, 1, value)
        cell.font = Font(name="Calibri", family=2, size=size, color="000000", bold=size > 11, underline="single" if row == 1 else None)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for row in (2, 3, 53, 59, 83, 88):
        sheet.row_dimensions[row].height = 19.5
    sheet.row_dimensions[1].height = 28.5
    sheet.row_dimensions[37].height = sheet.row_dimensions[45].height = 21


def write_sample_workbook(path: Path, sample: str, records: list[dict[str, Any]], *,
                          run_date: str, imgt_version: str, imgt_release: str,
                          lymphotrack_version: str = "") -> Path:
    """Write a local workbook with factual results and blank review areas."""
    path = Path(path).expanduser().resolve()
    ensure_outside_git(path)
    if path.exists():
        raise FileExistsError(f"Output already exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    overview = wb.active
    overview.title = "Oversikt"
    _configure(overview)
    _overview(overview, sample, records, run_date, imgt_version, imgt_release, lymphotrack_version)
    for rec in records:
        title = re.sub(r"[\\/*?:\[\]]", "_", rec["rank_label"])[:31]
        sheet = wb.create_sheet(title)
        _configure(sheet)
        _rearrangement(sheet, sample, rec)
    _sanger(wb.create_sheet("Mal for Sanger"), sample)
    for sheet in wb:
        if not sheet.print_area:
            sheet.print_area = f"A1:{'A' if sheet.title == 'Mal for Sanger' else 'D'}{sheet.max_row}"
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        wb.save(temporary)
        # Publish the complete file atomically; a competing destination is never replaced.
        os.link(temporary, path)
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)
    return path
