from __future__ import annotations

import os
import tempfile
from collections.abc import Iterable
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import MergedRow
from .privacy import ensure_outside_git

HEADERS = (
    "External ID",
    "Rank",
    "Sequence",
    "Length",
    "Merge count",
    "V-gene",
    "J-gene",
    "% total reads",
    "Cumulative %",
    "Mutation rate to partial V-gene (%)",
    "In-frame (Y/N)",
    "No Stop codon (Y/N)",
    "V-coverage",
    "CDR3 Seq",
    "Identity to VH",
    "Sample",
    "Reads",
    "Target",
    "Molecule",
    "Date PCR setup",
    "find sequence in other samples",
    "Subset",
    "Comment",
    "Fasta",
)


class ExcelReportWriter:
    def write(
        self,
        rows: tuple[MergedRow, ...] | list[MergedRow],
        output: Path,
        *,
        overwrite: bool = False,
        highlight_rows: Iterable[int] = (),
    ) -> Path:
        output = output.expanduser().resolve()
        if output.suffix.lower() != ".xlsx":
            output = output.with_suffix(".xlsx")
        ensure_outside_git(output)
        if output.exists() and not overwrite:
            raise FileExistsError(f"Output finnes allerede: {output}")
        output.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        worksheet.append(HEADERS)
        for excel_row, item in enumerate(rows, start=2):
            source = item.source
            worksheet.append(
                [
                    item.external_id,
                    source.rank,
                    source.sequence,
                    source.length,
                    source.merge_count,
                    source.v_gene,
                    source.j_gene,
                    source.percent_total_reads,
                    source.cumulative_percent,
                    source.mutation_rate,
                    source.in_frame,
                    source.no_stop_codon,
                    source.v_coverage,
                    source.cdr3_sequence,
                    f'=IF(J{excel_row}="","",100-J{excel_row})',
                    item.sample,
                    item.total_reads,
                    item.target,
                    item.molecule_type,
                    item.run_date,
                    item.other_samples,
                    item.subset,
                    item.comment,
                    item.fasta,
                ]
            )

        self._format(worksheet, len(rows) + 1, frozenset(highlight_rows))
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

        fd, temporary_name = tempfile.mkstemp(
            prefix=f".{output.stem}-", suffix=".xlsx", dir=output.parent
        )
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            os.replace(temporary, output)
        finally:
            workbook.close()
            temporary.unlink(missing_ok=True)
        return output

    def _format(
        self, worksheet, last_row: int, highlight_rows: frozenset[int]
    ) -> None:
        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in worksheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="000000")
            cell.alignment = Alignment(horizontal="center", vertical="bottom")
            cell.border = border
        worksheet.row_dimensions[1].height = 28

        yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")
        for result_index, row in enumerate(
            worksheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=24)
        ):
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="000000")
                cell.alignment = Alignment(vertical="center")
                cell.border = border
                if result_index in highlight_rows:
                    cell.fill = yellow_fill
            row[0].font = Font(name="Calibri", size=11, bold=True, color="1F2937")
            row[15].font = Font(name="Calibri", size=11, bold=True, color="000000")
            # Keep compact rows as in reference file. FASTA content can
            # be read in formula bar without entire sequence blowing up row height.
            row[23].alignment = Alignment(vertical="center", wrap_text=False)

        for column in ("H", "I", "J", "M", "O"):
            for cell in worksheet[f"{column}2:{column}{last_row}"]:
                cell[0].number_format = "0.##"
        for cell in worksheet[f"Q2:Q{last_row}"]:
            cell[0].number_format = "#,##0"

        green_font = Font(color="008000", bold=False)
        black_font = Font(color="000000", bold=False)
        bold_font = Font(color="000000", bold=True)
        red_font = Font(color="FF0000")
        percent_range = f"H2:H{last_row}"
        worksheet.conditional_formatting.add(
            percent_range, CellIsRule(operator="greaterThanOrEqual", formula=["2.5"], font=bold_font)
        )
        worksheet.conditional_formatting.add(
            percent_range,
            CellIsRule(operator="between", formula=["2.4", "2.499999999"], font=black_font),
        )
        worksheet.conditional_formatting.add(
            percent_range, CellIsRule(operator="lessThan", formula=["2.4"], font=green_font)
        )
        worksheet.conditional_formatting.add(
            f"K2:K{last_row}",
            FormulaRule(formula=['OR(K2="N",K2="N/A")'], font=red_font),
        )
        worksheet.conditional_formatting.add(
            f"L2:L{last_row}",
            FormulaRule(formula=['OR(L2="N",L2="N/A")'], font=red_font),
        )

        widths = {
            "A": 16, "B": 8, "C": 55, "D": 10, "E": 14, "F": 20, "G": 16,
            "H": 14, "I": 14, "J": 30, "K": 16, "L": 20, "M": 13,
            "N": 22, "O": 17, "P": 18, "Q": 14, "R": 11, "S": 12,
            "T": 18, "U": 34, "V": 13, "W": 30, "X": 60,
        }
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:X{last_row}"
        worksheet.sheet_view.showGridLines = True
