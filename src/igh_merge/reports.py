from __future__ import annotations

import json
import re
import hashlib
import shutil
from html import escape
from dataclasses import asdict, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, Image, KeepTogether, CondPageBreak

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .external import ArrestBatchResult, ExternalBatch, ImgtBatchResult, ImgtRecord, sequence_sha256
from .models import MergedRow
from .privacy import ensure_outside_git
from .evidence import EvidenceManifest, validate_evidence_manifest
from .excel import ExcelReportWriter
from .imgt_observations import j_call_notes, record_observation_text
from .matching import CLONE_THRESHOLD, exact_supporting_leaders
from .report_workbook import write_sample_workbook
from .edge_cdp import IMGT_SCREENSHOT_SPECS


REPORT_IMAGE_SECTIONS = (
    '01_v_gene', '02_d_gene', '03_j_gene', '06_junction', '09_v_region_translation',
)


def _gene(value: str) -> str:
    match = re.search(r"(IGH[VDJ][A-Za-z0-9/-]+\*\d+)", value)
    return match.group(1) if match else value


def _cdr3_length(record: ImgtRecord) -> str:
    raw = record.raw_cdr3_imgt_length
    derived = record.cdr3_aa_length
    if raw:
        if derived is not None and (not raw.isdecimal() or int(raw) != derived):
            return f'IMGT: {raw}; returned AA: {derived}'
        return f'{raw} (IMGT)'
    if derived is not None:
        return f'{derived} (returned AA; IMGT length not supplied)'
    return 'Not supplied'


def _display_observations(record: ImgtRecord) -> str:
    text = record_observation_text(record)
    text = re.sub(r'\bmanual assessment\s*;?\s*', '', text)
    text = text.replace('IMGT 4_IMGT-gapped-AA-sequences.txt ', '')
    return re.sub(r';\s*', ';\n', text)


def _summary_rows(record: ImgtRecord, summary: str) -> tuple[tuple[str, ...], ...]:
    lines = [line for line in summary.splitlines() if line.strip()]
    if any('\t' in line for line in lines):
        return tuple(tuple(line.split('\t')) for line in lines)
    fields = (
        ('External ID', record.external_id),
        ('V-DOMAIN Functionality', record.functionality),
        ('V-GENE and allele', record.v_call),
        ('V-REGION identity %', record.raw_identity_percent or record.v_identity_percent),
        ('V-REGION identity nt', f'{record.v_identity_numerator}/{record.v_identity_denominator}'
         if record.v_identity_numerator is not None else ''),
        ('V-REGION identity % (with ins/del events)', record.raw_identity_with_indel_events),
        ('J-GENE and allele', record.j_call),
        ('D-GENE and allele', record.d_call),
        ('CDR3-IMGT length', record.raw_cdr3_imgt_length),
        ('CDR3-IMGT (AA)', record.cdr3_aa),
        ('AA JUNCTION', record.junction_aa),
        ('JUNCTION frame', record.junction_frame),
        ('Insertions', record.insertions), ('Deletions', record.deletions),
    )
    return tuple((label, str(value)) for label, value in fields if value is not None and value != '')


def _sample_report_records(
    rows: tuple[MergedRow, ...] | list[MergedRow],
    batch: ExternalBatch,
    imgt: ImgtBatchResult,
    arrest: ArrestBatchResult | None,
    evidence: EvidenceManifest | None = None,
) -> dict[str, list[dict[str, Any]]]:
    """Group rows by sample and produce a serialisable record for each
    rearrangement (used by both the PDF and Word renderers)."""
    candidates = batch.by_id()
    evidence_by_id = {e.external_id: e for e in evidence.entries} if evidence else {}
    imgt_records = {item.external_id: item for item in imgt.records}
    arrest_records = (
        {item.external_id: item for item in arrest.records} if arrest else {}
    )
    # Validate every mapping before deciding how it is presented. Supporting
    # FR1 results stay in the complete analysis and audit, even without a tab.
    for candidate in batch.candidates:
        local = rows[candidate.row_index]
        if (local.sample != candidate.sample or local.target != candidate.target
            or local.source.rank != candidate.rank or local.molecule_type != candidate.molecule_type
            or sequence_sha256(local.source.sequence) != candidate.sequence_sha256
            or sequence_sha256(candidate.sequence) != candidate.sequence_sha256):
            raise ValueError('Report row does not match candidate ID/sequence/target/rank/molecule')
    by_sample: dict[str, list[str]] = {}
    for external_id in batch.candidate_ids:
        by_sample.setdefault(
            candidates[external_id].sample, []
        ).append(external_id)
    result: dict[str, list[dict[str, Any]]] = {}
    for sample, external_ids in by_sample.items():
        local_rows = [r for r in rows if r.sample == sample]
        depths = {
            target: next(
                (r.total_reads for r in local_rows if r.target == target), None
            )
            for target in ("Leader", "FR1")
        }
        selected_leaders = [
            rows[candidates[external_id].row_index]
            for external_id in external_ids
            if candidates[external_id].target == 'Leader'
            and rows[candidates[external_id].row_index].source.percent_total_reads >= CLONE_THRESHOLD
        ]
        selected_ids = {id(rows[candidates[i].row_index]): i for i in external_ids}
        supporting_by_leader: dict[str, list[str]] = {}
        supporting_ids: set[str] = set()
        for external_id in external_ids:
            local = rows[candidates[external_id].row_index]
            for leader in exact_supporting_leaders(local, selected_leaders):
                leader_id = selected_ids[id(leader)]
                supporting_by_leader.setdefault(leader_id, []).append(external_id)
                supporting_ids.add(external_id)
        ordered_ids = sorted(
            (external_id for external_id in external_ids if external_id not in supporting_ids),
            key=lambda external_id: (
                candidates[external_id].target != 'Leader',
                candidates[external_id].rank,
                candidates[external_id].row_index,
            ),
        )
        sample_records: list[dict[str, Any]] = []
        for external_id in ordered_ids:
            candidate = candidates[external_id]
            local = rows[candidate.row_index]
            record = imgt_records.get(external_id)
            missing_imgt = record is None
            if record is None:
                record = ImgtRecord(external_id, None, None, None, '', '', '', None, '', '', '', '', '')
            entry = evidence_by_id.get(external_id)
            arrest_record = arrest_records.get(external_id)
            support_ids = sorted(
                supporting_by_leader.get(external_id, []),
                key=lambda item: (candidates[item].rank, candidates[item].row_index),
            )
            fr1_support = [
                row.source.percent_total_reads
                for row in sorted(local_rows, key=lambda row: row.source.rank)
                if local.target == 'Leader' and exact_supporting_leaders(row, [local])
            ]
            identity = (
                f"{record.selected_identity_percent:.1f} %"
                if record.selected_identity_percent is not None
                else "not calculated"
            )
            counts = (
                f"{record.selected_identity_counts[0]}/{record.selected_identity_counts[1]} nt"
                if record.selected_identity_counts[0] is not None
                else "not available"
            )
            functionality = (
                "Functional (productive)"
                if record.productive is True
                else "Non-functional (unproductive)"
                if record.productive is False
                else record.functionality or "Not conclusively determined"
            )
            unproductive = record.productive is False or (
                record.productive is None
                and (local.source.in_frame == 'N' or local.source.no_stop_codon == 'N')
            )
            rank_label = f'{local.target}-{local.source.rank}'
            if unproductive:
                rank_label += ' (Ufunksjonell)'
            subset = (
                (arrest_record.subset if arrest_record else "")
                or record.cll_subset
                or "Not assigned / not analyzed"
            )
            comment = local.comment or ''
            if local.target == 'Leader' and local.source.percent_total_reads < CLONE_THRESHOLD:
                comment = ' | '.join(filter(None, [comment,
                    f'Manually selected below {CLONE_THRESHOLD:g}% '
                    f'({local.source.percent_total_reads:.2f} %); specialist review required']))
            arrest_metadata = asdict(arrest_record) if arrest_record else {}
            if arrest_record and arrest:
                arrest_metadata.update(results_url=arrest.results_url, endpoint=arrest.endpoint)
            arrest_notes = '; '.join(
                f'{label}: {value}' for label, value in (
                    ('SeqCure', arrest_record.seqcure if arrest_record else ''),
                    ('Subset', arrest_record.subset if arrest_record else ''),
                    ('Nearest subset', arrest_record.nearest_subset if arrest_record else ''),
                    ('Confidence', arrest_record.confidence if arrest_record else ''),
                    ('Score', arrest_record.score if arrest_record else ''),
                ) if value
            ) or ('No notes provided' if arrest_record else 'ARResT analysis missing')
            sample_records.append(
                {
                    "number": str(len(sample_records) + 1),
                    "rank_label": rank_label,
                    "sequence": candidate.sequence,
                    "fasta": f'>{candidate.sample}-{candidate.target}-{candidate.rank}\n{candidate.sequence}',
                    "supporting_external_ids": support_ids,
                    "leader_fraction": f"{local.source.percent_total_reads:.2f} %",
                    "fr1_support": (
                        ", ".join(
                            f"{value:.2f} %" for value in fr1_support
                        )
                        or "Not detected"
                    ),
                    "genes": " / ".join(
                        [
                            _gene(record.v_call),
                            _gene(record.d_call) or "-",
                            record.j_call,
                        ]
                    ),
                    "identity": f"{identity} ({counts})",
                    "functionality": functionality,
                    "cdr3": record.cdr3_aa or (
                        record.junction_aa + ' (AA JUNCTION; CDR3 not supplied)'
                        if record.junction_aa else '-'
                    ),
                    "cdr3_length": _cdr3_length(record),
                    "raw_cdr3_imgt_length": record.raw_cdr3_imgt_length,
                    "cdr3_aa_length": record.cdr3_aa_length,
                    "indel": 'Not analyzed / IMGT result missing' if missing_imgt else (
                        "; ".join(filter(None, [record.insertions, record.deletions])) or 'Not detected'
                    ),
                    "subset": subset,
                    "imgt_subset": record.cll_subset or 'Not assigned / not analyzed',
                    "arrest_subset": arrest_record.subset if arrest_record else 'ARResT analysis missing',
                    "arrest_notes": arrest_notes,
                    "arrest_seqcure_details": '\n'.join(getattr(arrest_record, 'seqcure_messages', ())),
                    "arrest_seqcure_advice": '\n'.join(getattr(arrest_record, 'seqcure_advice', ())),
                    "arrest_seqcure_status": getattr(arrest_record, 'seqcure_details_status', '') or 'Not retrieved',
                    "arrest_metadata": arrest_metadata,
                    "identity_source": record.selected_identity_source,
                    "raw_identities": f'{record.raw_identity_percent or "-"}; with indel events: {record.raw_identity_with_indel_events or "-"}',
                    "observations": _display_observations(record),
                    "j_notes": j_call_notes(record.j_call) or '-',
                    "analysis_status": 'IMGT analysis missing' if missing_imgt else 'IMGT result available',
                    "summary_text": entry.sections.get('00_summary', '') if entry else 'IMGT evidence missing',
                    "summary_rows": _summary_rows(record, entry.sections.get('00_summary', '') if entry else ''),
                    "images": entry.images if entry else (),
                    "missing_evidence": '; '.join(f'{k}: {v}' for k,v in entry.missing_sections.items()) if entry else 'IMGT evidence missing',
                    "sequence_sha256": candidate.sequence_sha256,
                    "session_id": batch.session_id,
                    "molecule_type": candidate.molecule_type,
                    "external_id": record.external_id,
                    "depth_leader": str(depths.get("Leader") or "-"),
                    "depth_fr1": str(depths.get("FR1") or "-"),
                    "comment": comment,
                }
            )
        result[sample] = sample_records
    return result


def _result_fields(rec: dict[str, Any]) -> list[tuple[str, str]]:
    return [
        ('External ID', rec['external_id']), ('IGHV / IGHD / IGHJ', rec['genes']),
        ('Target fraction', rec['leader_fraction']), ('FR1 support', rec['fr1_support']),
        ('VH identity', rec['identity']), ('Functionality', rec['functionality']),
        ('CDR3 (AA)', rec['cdr3']), ('CDR3 length', rec['cdr3_length']),
        ('Insertion / deletion', rec['indel']), ('IMGT subset', rec['imgt_subset']),
        ('ARResT subset', rec['arrest_subset']), ('ARResT notes', rec['arrest_notes']),
        ('ARResT SeqCure details', rec['arrest_seqcure_details'] or rec['arrest_seqcure_status']),
        ('ARResT SeqCure advice', rec['arrest_seqcure_advice'] or '-'),
        ('Raw identities', rec['raw_identities']), ('J call notes', rec['j_notes']),
        ('Manual observations', rec['observations']), ('Analysis status', rec['analysis_status']),
        ('Molecule type', rec['molecule_type']), ('Comment', rec['comment'] or '-'),
    ]


def _metadata_items(sample, run_date, imgt_version, imgt_release, has_arrest, lymphotrack_version):
    return (
        ('Sample', sample), ('Run date', run_date),
        ('Analysis', 'LymphoTrack IGH SHM (HTS)'),
        ('Analysis tools', 'LymphoTrack, IMGT/V-QUEST' + (' and ARResT/AssignSubsets' if has_arrest else '')),
        ('LymphoTrack version', lymphotrack_version or 'not provided'),
        ('IMGT/V-QUEST version', imgt_version or 'not provided'),
        ('IMGT reference release', imgt_release or 'not provided'),
    )


def _report_images(rec):
    by_section = {Path(filename).stem: filename for filename in rec['images']}
    return [by_section[key] for key in REPORT_IMAGE_SECTIONS if key in by_section]


def _image_dimensions(filename):
    try:
        width, height = ImageReader(filename).getSize()
        return (width, height) if width > 0 and height > 0 else None
    except (OSError, ValueError):
        return None


def _write_sample_pdf(
    path: Path,
    sample: str,
    records: list[dict[str, Any]],
    run_date: str,
    imgt_version: str,
    imgt_release: str,
    has_arrest: bool,
    lymphotrack_version: str = '',
) -> None:
    styles = getSampleStyleSheet()
    for style_name in ('Title', 'Heading2', 'Heading3', 'BodyText'):
        styles[style_name].textColor = colors.black
    styles['Title'].fontName = 'Helvetica-Bold'
    styles['Title'].fontSize = 18
    styles['Heading2'].fontSize = 13
    styles['Heading2'].spaceBefore = 10
    styles['Heading2'].spaceAfter = 6
    styles['BodyText'].fontSize = 9.5
    styles['BodyText'].leading = 12
    sequence_style = ParagraphStyle('Sequence', parent=styles['BodyText'], fontName='Courier', fontSize=8, leading=10)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm,
                            rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm)
    width = doc.width - 12  # Account for the frame's default inner padding.
    story = [Paragraph('IGHV SHM report draft', styles['Title'])]
    for label, value in _metadata_items(sample, run_date, imgt_version, imgt_release, has_arrest, lymphotrack_version):
        story.append(Paragraph(f'<b>{escape(label)}:</b> {escape(value)}', styles['BodyText']))
    if records:
        first = records[0]
        story.append(Paragraph(f"<b>Depth:</b> {first['depth_leader']} reads (Leader), "
                               f"{first['depth_fr1']} reads (FR1)", styles['BodyText']))

    def table_for(fields):
        data = [['Field', 'Result']] + [[str(label), str(value)] for label, value in fields]
        header_style = ParagraphStyle('TableHeader', parent=styles['BodyText'], textColor=colors.white, fontName='Helvetica-Bold')
        data = [[Paragraph(escape(value).replace('\n', '<br/>'), header_style if index == 0 else styles['BodyText'])
                 for value in row] for index, row in enumerate(data)]
        table = Table(data, colWidths=[44 * mm, width - 44 * mm], repeatRows=1, splitByRow=1, splitInRow=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17334A')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6F8')]),
            ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#D9D9D9')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        return table

    for rec in records:
        story.extend([CondPageBreak(40 * mm), Spacer(1, 4 * mm),
                      Paragraph(f"Rearrangement {rec['number']} - {escape(rec['rank_label'])}", styles['Heading2'])])
        header, _, sequence = rec['fasta'].partition('\n')
        wrapped = '<br/>'.join(escape(sequence[i:i + 78]) for i in range(0, len(sequence), 78))
        story.extend([Paragraph(escape(header), sequence_style), Paragraph(wrapped, sequence_style), Spacer(1, 3 * mm),
                      table_for(_result_fields(rec)), Paragraph('IMGT result fields', styles['Heading2'])])
        summary_fields = [(row[0], '\n'.join(row[1:])) for row in rec['summary_rows']]
        if summary_fields:
            story.append(table_for(summary_fields))
        else:
            story.append(Paragraph(escape(rec['summary_text']), styles['BodyText']))
        if summary_fields and '\t' not in rec['summary_text'] and rec['summary_text'] != 'IMGT evidence missing':
            story.append(Paragraph('Captured IMGT summary', styles['Heading3']))
            for line in rec['summary_text'].splitlines():
                if line.strip():
                    story.append(Paragraph(escape(line), styles['BodyText']))
        if rec['missing_evidence']:
            story.append(Paragraph(escape(rec['missing_evidence']), styles['BodyText']))
        for filename in _report_images(rec):
            dimensions = _image_dimensions(filename)
            caption = Path(filename).stem.replace('_', ' ')
            if dimensions is None:
                story.append(Paragraph(f'Image unavailable: {escape(caption)}', styles['BodyText']))
                continue
            image = Image(filename)
            factor = min(width / dimensions[0], 210 * mm / dimensions[1], 1)
            image.drawWidth, image.drawHeight = dimensions[0] * factor, dimensions[1] * factor
            story.append(KeepTogether([Spacer(1, 4 * mm), Paragraph(escape(caption), styles['Heading3']), image]))
        story.extend([Spacer(1, 5 * mm), Paragraph('Specialist assessment: __________________________', styles['BodyText'])])
    story.extend([Spacer(1, 5 * mm), Paragraph('DRAFT - requires specialist approval before clinical use.', styles['BodyText'])])

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#526371'))
        canvas.drawRightString(A4[0] - 18 * mm, 9 * mm, f'DRAFT  |  {document.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def _write_sample_docx(
    path: Path,
    sample: str,
    records: list[dict[str, Any]],
    run_date: str,
    imgt_version: str,
    imgt_release: str,
    has_arrest: bool,
    lymphotrack_version: str = '',
) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(16)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    for name, size in (('Title', 18), ('Heading 2', 13), ('Heading 3', 11)):
        style = document.styles[name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
    document.add_paragraph('IGHV SHM report draft', 'Title')
    for label, value in _metadata_items(sample, run_date, imgt_version, imgt_release, has_arrest, lymphotrack_version):
        paragraph = document.add_paragraph()
        paragraph.add_run(label + ': ').bold = True
        paragraph.add_run(value)
    if records:
        first = records[0]
        paragraph = document.add_paragraph()
        paragraph.add_run('Depth: ').bold = True
        paragraph.add_run(f"{first['depth_leader']} reads (Leader), {first['depth_fr1']} reads (FR1)")

    def table_for(fields):
        table = document.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for column, width in zip(table.columns, (Mm(44), Mm(124))):
            column.width = width
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            border = OxmlElement('w:' + edge)
            for attr, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                border.set(qn('w:' + attr), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        for index, (label, value) in enumerate([('Field', 'Result')] + list(fields)):
            row = table.add_row()
            for column, text in enumerate((str(label), str(value))):
                cell = row.cells[column]
                cell.width = Mm(44 if column == 0 else 124)
                cell.text = text
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), '17334A' if index == 0 else ('F3F6F8' if index % 2 == 0 else 'FFFFFF'))
                cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0 or column == 0
                        run.font.color.rgb = RGBColor(255, 255, 255) if index == 0 else RGBColor(0, 0, 0)
            if index == 0:
                repeated = OxmlElement('w:tblHeader')
                row._tr.get_or_add_trPr().append(repeated)
        return table

    for rec in records:
        document.add_heading(f"Rearrangement {rec['number']} - {rec['rank_label']}", level=2)
        header, _, sequence = rec['fasta'].partition('\n')
        paragraph = document.add_paragraph()
        run = paragraph.add_run(header + '\n' + '\n'.join(sequence[i:i + 78] for i in range(0, len(sequence), 78)))
        run.font.name, run.font.size = 'Courier New', Pt(9)
        table_for(_result_fields(rec))
        document.add_heading('IMGT result fields', level=2)
        summary_fields = [(row[0], '\n'.join(row[1:])) for row in rec['summary_rows']]
        if summary_fields:
            table_for(summary_fields)
        else:
            document.add_paragraph(rec['summary_text'])
        if summary_fields and '\t' not in rec['summary_text'] and rec['summary_text'] != 'IMGT evidence missing':
            document.add_heading('Captured IMGT summary', level=3)
            document.add_paragraph(rec['summary_text'])
        if rec['missing_evidence']:
            document.add_paragraph(rec['missing_evidence'])
        for filename in _report_images(rec):
            dimensions = _image_dimensions(filename)
            caption = Path(filename).stem.replace('_', ' ')
            if dimensions is None:
                document.add_paragraph(f'Image unavailable: {caption}')
                continue
            document.add_heading(caption, level=3)
            width, height = dimensions
            factor = min(168 / width, 210 / height, 25.4 / 96)
            document.add_picture(filename, width=Mm(width * factor), height=Mm(height * factor))
        document.add_paragraph('Specialist assessment: __________________________')
    document.add_paragraph('DRAFT - requires specialist approval before clinical use.')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('DRAFT  |  ')
    page_number = OxmlElement('w:fldSimple')
    page_number.set(qn('w:instr'), 'PAGE')
    footer._p.append(page_number)
    document.save(str(path))


def generate_clinical_report_package(
    run_directory: Path,
    rows: tuple[MergedRow, ...] | list[MergedRow],
    batch: ExternalBatch,
    imgt: ImgtBatchResult,
    arrest: ArrestBatchResult | None = None,
    *, evidence: EvidenceManifest | None = None,
    lymphotrack_version: str = '',
) -> Path:
    output = run_directory.resolve() / f"{batch.run_date}_reports" / batch.session_id
    ensure_outside_git(output)
    if evidence:
        validate_evidence_manifest(evidence, batch)
    if output.exists():
        output = output / datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    output.mkdir(parents=True, exist_ok=False)
    by_sample = _sample_report_records(rows, batch, imgt, arrest, evidence)
    export_rows = [replace(row) for row in rows]
    for candidate in batch.candidates:
        export_rows[candidate.row_index].external_id = candidate.external_id
    ExcelReportWriter().write(export_rows, output / 'analysis.xlsx')
    if evidence:
        evidence_directory = output / 'evidence'
        evidence_directory.mkdir()
        for entry in evidence.entries:
            directory = evidence_directory / entry.external_id
            directory.mkdir()
            (directory / 'summary.txt').write_text(entry.sections.get('00_summary', ''), encoding='utf-8')
            packaged_entry = asdict(entry)
            packaged_entry['images'] = [Path(filename).name for filename in entry.images]
            packaged_entry['image_hashes'] = {
                Path(filename).name: entry.image_hashes[filename] for filename in entry.images
            }
            (directory / 'evidence.audit.json').write_text(json.dumps(packaged_entry, ensure_ascii=False, indent=2), encoding='utf-8')
            for filename in entry.images:
                shutil.copy2(filename, directory / Path(filename).name)
    safe_names = {
        sample: re.sub(r"[^A-Za-z0-9_.-]", "_", sample)
        for sample in by_sample
    }
    counts = {}
    for sample, name in list(safe_names.items()):
        key = name.casefold()
        counts[key] = counts.get(key, 0) + 1
    for sample, name in list(safe_names.items()):
        if counts[name.casefold()] > 1:
            safe_names[sample] = name + '_' + hashlib.sha256(sample.encode()).hexdigest()[:8]
    for sample, records in by_sample.items():
        name = safe_names[sample]
        _write_sample_docx(
            output / f"{name}_IGHV_report_draft.docx",
            sample,
            records,
            batch.run_date,
            imgt.program_version,
            imgt.reference_release,
            has_arrest=bool(arrest),
            lymphotrack_version=lymphotrack_version,
        )
        write_sample_workbook(
            output / f'{name}_IGHV_analysis.xlsx', sample, records,
            run_date=batch.run_date, imgt_version=imgt.program_version,
            imgt_release=imgt.reference_release, lymphotrack_version=lymphotrack_version,
        )
    imgt_ids = {record.external_id for record in imgt.records}
    arrest_ids = {record.external_id for record in arrest.records} if arrest else set()
    evidence_by_id = {entry.external_id: entry for entry in evidence.entries} if evidence else {}
    required_sections = [spec.key for spec in IMGT_SCREENSHOT_SPECS if not spec.optional]
    missing_evidence_ids = [
        external_id for external_id in batch.candidate_ids
        if external_id not in evidence_by_id
    ]
    missing_evidence_sections_by_id = {}
    for external_id in batch.candidate_ids:
        entry = evidence_by_id.get(external_id)
        if entry is None:
            continue
        image_sections = {Path(filename).stem for filename in entry.images}
        missing_sections = [
            section for section in required_sections
            if not entry.sections.get(section, '').strip()
            or section not in image_sections
            or section in entry.missing_sections
        ]
        if missing_sections:
            missing_evidence_sections_by_id[external_id] = missing_sections
    reported_ids = [record['external_id'] for records in by_sample.values() for record in records]
    supporting_by_reported_id = {
        record['external_id']: record['supporting_external_ids']
        for records in by_sample.values() for record in records if record['supporting_external_ids']
    }
    supporting_ids = {external_id for ids in supporting_by_reported_id.values() for external_id in ids}
    imgt_complete = set(batch.candidate_ids) == imgt_ids
    audit = {
        "schema_version": 1,
        "complete": imgt_complete,
        "imgt_complete": imgt_complete,
        "missing_imgt_ids": [i for i in batch.candidate_ids if i not in imgt_ids],
        "arrest_complete": set(batch.candidate_ids) == arrest_ids,
        "missing_arrest_ids": [i for i in batch.candidate_ids if i not in arrest_ids],
        "evidence_coverage_complete": not missing_evidence_ids,
        "evidence_complete": not missing_evidence_ids and not missing_evidence_sections_by_id,
        "missing_evidence_ids": missing_evidence_ids,
        "missing_evidence_sections_by_id": missing_evidence_sections_by_id,
        "reported_ids": reported_ids,
        "supporting_ids": [i for i in batch.candidate_ids if i in supporting_ids],
        "supporting_by_reported_id": supporting_by_reported_id,
        "evidence_ids": [e.external_id for e in evidence.entries] if evidence else [],
        "mapping": [asdict(c) for c in batch.candidates],
        "status": "DRAFT - requires specialist approval",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "run_date": batch.run_date,
        "session_id": batch.session_id,
        "imgt_version": imgt.program_version,
        "imgt_reference_release": imgt.reference_release,
        "lymphotrack_version": lymphotrack_version,
        "imgt_records": [asdict(item) for item in imgt.records],
        "arrest_records": [asdict(item) for item in arrest.records]
        if arrest
        else [],
    }
    (output / "report_data.audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return output
