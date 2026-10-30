"""Local, portable run snapshots and recovery of existing analysis audits."""
from dataclasses import asdict, dataclass, replace
from pathlib import Path, PurePosixPath
import hashlib
import json
import os
import re
import tempfile
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED, ZIP_STORED

from .models import MergeResult, MergedRow, RunManifest, SampleFile, SourceRow
from .qc import ControlLimits, evaluate_qc
from .service import MergeService
from .external import (
    ExternalBatch, ExternalCandidate, ImgtRecord, ImgtBatchResult, ArrestRecord,
    ArrestBatchResult, apply_imgt_result, apply_arrest_result, sequence_sha256,
)
from .evidence import EvidenceEntry, EvidenceManifest, validate_evidence_manifest
from .privacy import ensure_outside_git
from .matching import annotate_exact_matches
from .io import SUMMARY_SUFFIX


def _merged_result(path: Path, run_date: str, limits: ControlLimits | None) -> MergeResult:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = workbook.active.iter_rows(values_only=True)
        headers = {str(value).strip().casefold(): i for i, value in enumerate(next(values))}
        def cell(row, name, default=None):
            index = headers.get(name.casefold())
            return row[index] if index is not None and index < len(row) else default
        required = ('Rank', 'Sequence', 'Sample', 'Target', 'Reads', '% total reads', 'Merge count')
        if any(name.casefold() not in headers for name in required):
            raise ValueError('Merged workbook does not contain the required source columns')
        def number(value, optional=False):
            if optional and (value is None or str(value).strip().lower() in {'', 'n/a', 'na', '-'}):
                return None
            return float(str(value).replace(',', '.'))
        rows, samples = [], {}
        for values_row in values:
            if not cell(values_row, 'Sequence'):
                continue
            sample, target = str(cell(values_row, 'Sample')), str(cell(values_row, 'Target'))
            if target not in {'Leader', 'FR1'}:
                raise ValueError('Unknown target in merged workbook')
            samples.setdefault(sample, len(samples)+1)
            sequence = str(cell(values_row, 'Sequence'))
            source = SourceRow(
                int(number(cell(values_row, 'Rank'))), sequence,
                int(number(cell(values_row, 'Length', len(sequence)))),
                int(number(cell(values_row, 'Merge count'))),
                str(cell(values_row, 'V-gene', '') or ''), str(cell(values_row, 'J-gene', '') or ''),
                number(cell(values_row, '% total reads')), number(cell(values_row, 'Cumulative %')),
                number(cell(values_row, 'Mutation rate to partial V-gene (%)'), True),
                str(cell(values_row, 'In-frame (Y/N)', '') or ''),
                str(cell(values_row, 'No Stop codon (Y/N)', '') or ''),
                number(cell(values_row, 'V-coverage'), True), str(cell(values_row, 'CDR3 Seq', '') or ''),
            )
            rows.append(MergedRow(source, sample, samples[sample], int(number(cell(values_row, 'Reads'))),
                                  target, str(cell(values_row, 'Molecule', 'gDNA') or 'gDNA'), run_date))
        if not rows:
            raise ValueError('Merged workbook contains no source rows')
        annotate_exact_matches(rows)
        files = []
        for sample in samples:
            for target in ('Leader', 'FR1'):
                group = [row for row in rows if row.sample == sample and row.target == target]
                if group:
                    first = group[0]
                    files.append(SampleFile(path, sample, first.sample_number, target, first.molecule_type,
                                            first.total_reads, tuple(row.source for row in group)))
        return MergeResult(RunManifest(path.parent, run_date, tuple(files), len(samples)), tuple(rows),
                           evaluate_qc(rows, limits))
    finally:
        workbook.close()


@dataclass(frozen=True)
class RunSession:
    result: MergeResult
    batch: ExternalBatch | None = None
    imgt: ImgtBatchResult | None = None
    arrest: ArrestBatchResult | None = None
    evidence: EvidenceManifest | None = None
    selected: tuple[int, ...] = ()
    lymphotrack_version: str = ''
    limits: ControlLimits = ControlLimits()


def _imgt(data):
    if not data:
        return None
    data = dict(data)
    records = []
    for row in data.pop('records'):
        row = dict(row)
        row['numbered_aa'] = tuple(tuple(pair) for pair in row.get('numbered_aa', ()))
        records.append(ImgtRecord(**row))
    return ImgtBatchResult(records=tuple(records), **data)


def _arrest(data):
    if not data:
        return None
    data = dict(data)
    records = []
    for row in data.pop('records'):
        row = dict(row)
        for key in ('seqcure_messages', 'seqcure_advice'):
            row[key] = tuple(row.get(key, ()))
        records.append(ArrestRecord(**row))
    return ArrestBatchResult(records=tuple(records), **data)


def validate_session(state: RunSession):
    if not re.fullmatch(r'\d{4}_\d{2}_\d{2}', state.result.manifest.run_date):
        raise ValueError('Invalid saved run date')
    if len(set(state.selected)) != len(state.selected) or any(type(i) is not int or not 0 <= i < len(state.result.rows) for i in state.selected):
        raise ValueError('Invalid saved sequence selection')
    if state.batch is None:
        if state.imgt or state.arrest or state.evidence:
            raise ValueError('Analysis has no saved sequence mapping')
        return
    batch = state.batch
    if not re.fullmatch(r'[A-Za-z0-9_-]+', batch.session_id) or batch.run_date != state.result.manifest.run_date:
        raise ValueError('Invalid saved analysis session')
    ids, indices = set(), set()
    for candidate in batch.candidates:
        if (not re.fullmatch(r'SEQ-[A-F0-9]{6,24}', candidate.external_id)
                or candidate.external_id in ids or candidate.row_index in indices
                or type(candidate.row_index) is not int or not 0 <= candidate.row_index < len(state.result.rows)):
            raise ValueError('Invalid saved sequence mapping')
        row = state.result.rows[candidate.row_index]
        if (candidate.sample, candidate.target, candidate.rank, candidate.molecule_type, candidate.sequence_sha256) != (
            row.sample, row.target, row.source.rank, row.molecule_type, sequence_sha256(row.source.sequence)
        ) or sequence_sha256(candidate.sequence) != candidate.sequence_sha256:
            raise ValueError('Saved analysis does not match source sequence')
        ids.add(candidate.external_id)
        indices.add(candidate.row_index)
    for result in (state.imgt, state.arrest):
        if result:
            observed = [record.external_id for record in result.records]
            if len(set(observed)) != len(observed) or not set(observed) <= ids:
                raise ValueError('Saved result does not match analysis mapping')
    if state.evidence:
        validate_evidence_manifest(state.evidence, batch)


def save_session(path: Path, state: RunSession) -> Path:
    path = path.expanduser().resolve()
    ensure_outside_git(path)
    validate_session(state)
    payload = asdict(state)
    payload['schema_version'] = 1
    assets = {}
    if state.evidence:
        for index, entry in enumerate(payload['evidence']['entries']):
            members, hashes = [], {}
            for filename in entry['images']:
                member = f'images/{index}/{Path(filename).name}'
                if member in assets:
                    raise ValueError('Duplicate saved image name')
                assets[member] = Path(filename)
                members.append(member)
                hashes[member] = entry['image_hashes'][filename]
            entry['images'], entry['image_hashes'] = members, hashes
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tempname = tempfile.mkstemp(dir=path.parent, suffix='.tmp')
    os.close(fd)
    try:
        with ZipFile(tempname, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('session.json', json.dumps(payload, ensure_ascii=False, default=str))
            for member, filename in assets.items():
                archive.write(filename, member, compress_type=ZIP_STORED)
        os.replace(tempname, path)
    finally:
        Path(tempname).unlink(missing_ok=True)
    return path


def load_session(path: Path, asset_directory: Path) -> RunSession:
    """Read JSON and bound images only; never execute or extract archive paths."""
    path = path.resolve()
    asset_directory = asset_directory.resolve()
    ensure_outside_git(asset_directory)
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            if len(set(names)) != len(names) or any(
                name != 'session.json' and not re.fullmatch(r'images/\d+/[A-Za-z0-9_-]+\.png', name)
                for name in names
            ):
                raise ValueError('Invalid saved run archive members')
            if sum(item.file_size for item in archive.infolist()) > 1024 * 1024 * 1024:
                raise ValueError('Saved run is too large')
            if any(item.file_size > 128 * 1024 * 1024 for item in archive.infolist()):
                raise ValueError('Saved run contains an oversized file')
            if archive.getinfo('session.json').file_size > 64 * 1024 * 1024:
                raise ValueError('Saved run metadata is too large')
            data = json.loads(archive.read('session.json'))
            if data.pop('schema_version') != 1:
                raise ValueError('Unsupported saved run version')
            raw = data.pop('result')
            manifest = raw['manifest']
            files = []
            for item in manifest['sample_files']:
                files.append(SampleFile(**{**item, 'path': Path(item['path']),
                                          'rows': tuple(SourceRow(**row) for row in item['rows'])}))
            manifest = RunManifest(**{**manifest, 'run_directory': path.parent,
                                      'sample_files': tuple(files), 'warnings': tuple(manifest['warnings'])})
            rows = tuple(MergedRow(**{**row, 'source': SourceRow(**row['source'])}) for row in raw['rows'])
            limits = ControlLimits(**data.pop('limits'))
            result = MergeResult(manifest, rows, evaluate_qc(list(rows), limits))
            batch_data = data.pop('batch')
            batch = ExternalBatch(**{**batch_data, 'candidates': tuple(ExternalCandidate(**c) for c in batch_data['candidates'])}) if batch_data else None
            evidence_data = data.pop('evidence')
            entries = []
            if evidence_data:
                for item in evidence_data['entries']:
                    images, hashes = [], {}
                    for member in item['images']:
                        if member not in names or member == 'session.json':
                            raise ValueError('Saved image is missing')
                        content = archive.read(member)
                        digest = hashlib.sha256(content).hexdigest()
                        if digest != item['image_hashes'].get(member):
                            raise ValueError('Saved image hash mismatch')
                        # Member has already passed the restrictive name check above.
                        target = asset_directory.joinpath(*PurePosixPath(member).parts)
                        target.parent.mkdir(parents=True, exist_ok=True)
                        target.write_bytes(content)
                        images.append(str(target))
                        hashes[str(target)] = digest
                    entries.append(EvidenceEntry(**{**item, 'images': tuple(images), 'image_hashes': hashes}))
            state = RunSession(result, batch, _imgt(data.pop('imgt')), _arrest(data.pop('arrest')),
                               EvidenceManifest(tuple(entries)) if evidence_data else None,
                               tuple(data.pop('selected')), data.pop('lymphotrack_version'), limits)
            validate_session(state)
            return state
    except (BadZipFile, KeyError, TypeError, AttributeError, json.JSONDecodeError) as exc:
        raise ValueError(f'Invalid saved run: {exc}') from exc


def restore_analysis(mapping_path: Path, limits: ControlLimits | None = None) -> RunSession:
    """Recover existing on-disk results without contacting IMGT or ARResT."""
    mapping_path = mapping_path.resolve()
    mapping = json.loads(mapping_path.read_text(encoding='utf-8'))
    if (not re.fullmatch(r'\d{4}_\d{2}_\d{2}', mapping.get('run_date', ''))
            or not re.fullmatch(r'[A-Za-z0-9_-]+', mapping.get('session_id', ''))):
        raise ValueError('Invalid saved analysis date or session')
    root = next((parent for parent in mapping_path.parents
                 if any(parent.glob('*_LEADER_output')) and any(parent.glob('*_IGH_FR1_output'))), None)
    if root is not None and any(root.glob(f'*_LEADER_output/*{SUMMARY_SUFFIX}')):
        result = MergeService().process(root, None, limits=limits)
    else:
        workbook = next((parent / name for parent in mapping_path.parents
                         for name in ('analysis.xlsx', f'{mapping["run_date"]}_merged.xlsx')
                         if (parent / name).is_file()), None)
        if workbook is None:
            raise ValueError('Original inputs or merged workbook are missing. Open a saved .ighv run instead.')
        result = _merged_result(workbook, mapping['run_date'], limits)
        # Outputs and old evidence belong to the selected analysis root, not an ancestor workbook folder.
        root = next((p for p in mapping_path.parents if (p / f'{mapping["run_date"]}_external_results').is_dir()), workbook.parent)
        result = replace(result, manifest=replace(result.manifest, run_directory=root))
    candidates = []
    for item in mapping.get('candidates', mapping.get('mapping', [])):
        matches = [(i, row) for i, row in enumerate(result.rows)
                   if (row.sample, row.target, row.source.rank, row.molecule_type, sequence_sha256(row.source.sequence)) ==
                   (item['sample'], item['target'], item['rank'], item.get('molecule_type', 'gDNA'), item['sequence_sha256'])]
        if len(matches) != 1:
            raise ValueError('Existing analysis does not match current input files')
        index, row = matches[0]
        candidate = ExternalCandidate(item['external_id'], index, row.sample, row.target, row.molecule_type,
                                      row.source.rank, row.source.sequence, item['sequence_sha256'])
        row.external_id, row.external_sequence_sha256 = candidate.external_id, candidate.sequence_sha256
        candidates.append(candidate)
    if not candidates:
        raise ValueError('No saved sequence mapping found')
    candidates.sort(key=lambda c: c.row_index)
    batch = ExternalBatch(mapping['session_id'], mapping['run_date'], mapping.get('created_at', ''), tuple(candidates))
    imgt_records, arrest_records, parameters = {}, {}, {}
    lymphotrack_version = mapping.get('lymphotrack_version', '')
    report_limits = mapping.get('control_limits')
    report_paths = sorted((root / f'{batch.run_date}_reports' / batch.session_id).glob('**/report_data.audit.json'))
    for report_path in report_paths:
        report = json.loads(report_path.read_text(encoding='utf-8'))
        if report.get('session_id') != batch.session_id:
            continue
        lymphotrack_version = report.get('lymphotrack_version') or lymphotrack_version
        report_limits = report.get('control_limits') or report_limits
        imgt_records.update({r['external_id']: r for r in report.get('imgt_records', [])})
        arrest_records.update({r['external_id']: r for r in report.get('arrest_records', [])})
        parameters.update({'IMGT/V-QUEST program version': report.get('imgt_version', ''),
                           'IMGT/V-QUEST reference directory release': report.get('imgt_reference_release', '')})
    if mapping_path.name == 'report_data.audit.json':
        imgt_records = {r['external_id']: r for r in mapping.get('imgt_records', [])}
        arrest_records = {r['external_id']: r for r in mapping.get('arrest_records', [])}
        parameters = {'IMGT/V-QUEST program version': mapping.get('imgt_version', ''),
                      'IMGT/V-QUEST reference directory release': mapping.get('imgt_reference_release', '')}
    else:
        for audit in sorted(mapping_path.parent.glob('*/results.audit.json')):
            saved = json.loads(audit.read_text(encoding='utf-8'))
            if saved.get('service') == 'IMGT/V-QUEST':
                imgt_records.update({r['external_id']: r for r in saved['records']})
                parameters.update(saved.get('parameters', {}))
            elif saved.get('service') == 'ARResT/AssignSubsets':
                arrest_records.update({r['external_id']: r for r in saved['records']})
    imgt = _imgt(dict(parameters=parameters, records=list(imgt_records.values()), raw_parameters='', raw_airr='')) if imgt_records else None
    arrest = _arrest(dict(records=list(arrest_records.values()), raw_tsv='', results_url='')) if arrest_records else None
    entries = {}
    audit_paths = list((root / f'{batch.run_date}_imgt_evidence' / batch.session_id).glob('**/evidence.audit.json'))
    audit_paths += list((root / f'{batch.run_date}_reports' / batch.session_id).glob('**/evidence.audit.json'))
    audit_paths += list(mapping_path.parent.glob('evidence/*/evidence.audit.json'))
    for audit in sorted(audit_paths):
        content = json.loads(audit.read_text(encoding='utf-8'))
        for item in content.get('entries', [content] if 'external_id' in content else []):
            images, hashes = [], {}
            for filename in item['images']:
                image = audit.parent / Path(filename).name
                if not image.is_file():
                    image = Path(filename)
                digest = item['image_hashes'][filename]
                if not image.is_file() or hashlib.sha256(image.read_bytes()).hexdigest() != digest:
                    raise ValueError('Existing IMGT image is missing or changed')
                images.append(str(image.resolve()))
                hashes[str(image.resolve())] = digest
            entries[item['external_id']] = EvidenceEntry(**{**item, 'images': tuple(images), 'image_hashes': hashes})
    effective_limits = ControlLimits(**report_limits) if report_limits else limits or ControlLimits()
    result = replace(result, qc=evaluate_qc(list(result.rows), effective_limits))
    state = RunSession(result, batch, imgt, arrest, EvidenceManifest(tuple(entries.values())) if entries else None,
                       tuple(c.row_index for c in candidates), lymphotrack_version, effective_limits)
    validate_session(state)
    if imgt:
        for record in imgt.records:
            part = replace(batch, candidates=(batch.by_id()[record.external_id],))
            apply_imgt_result(result.rows, part, replace(imgt, records=(record,)))
    if arrest:
        for record in arrest.records:
            part = replace(batch, candidates=(batch.by_id()[record.external_id],))
            apply_arrest_result(result.rows, part, replace(arrest, records=(record,)))
    return state
